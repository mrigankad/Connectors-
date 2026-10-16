"""Generate the DOCX/PPTX fixtures under fixtures/office. Re-run only when fixtures change; outputs are committed.

These are python-docx/python-pptx generated files, i.e. fixture-level evidence. Word/PowerPoint-authored
samples should be added alongside them (sanitized) before claiming anything beyond fixture_tested.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent.parent / "fixtures" / "office"


def _png_1x1() -> bytes:
    import struct
    import zlib

    def chunk(kind: bytes, body: bytes) -> bytes:
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF)

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    idat = zlib.compress(b"\x00\xff\x00\x00")
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


PNG_1X1 = _png_1x1()


def build_docx() -> None:
    d = Document()
    d.add_heading("Quarterly Review", level=1)
    d.add_paragraph("Revenue grew 12% quarter over quarter.")
    d.add_paragraph("")  # empty paragraph keeps its index
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Metric", "Value"
    t.cell(1, 0).text, t.cell(1, 1).text = "Churn", "2.1%"
    d.add_paragraph("Appendix follows.")
    d.save(ROOT / "review.docx")
    ROOT.joinpath("review.png").write_bytes(PNG_1X1)
    d2 = Document()
    d2.add_paragraph("Diagram below.")
    d2.add_picture(str(ROOT / "review.png"), width=Inches(1))
    d2.save(ROOT / "with_image.docx")
    ROOT.joinpath("review.png").unlink()


def build_pptx() -> None:
    p = Presentation()
    s1 = p.slides.add_slide(p.slide_layouts[1])
    s1.shapes.title.text = "Roadmap 2027"
    s1.placeholders[1].text_frame.text = "Launch SAMYAK pilot"
    s1.notes_slide.notes_text_frame.text = "Speaker: mention the bank partner."
    s2 = p.slides.add_slide(p.slide_layouts[5])
    s2.shapes.title.text = "Numbers"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "KPI", "Target"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "Lag", "< 5 min"
    s3 = p.slides.add_slide(p.slide_layouts[6])  # blank
    ROOT.joinpath("tmp.png").write_bytes(PNG_1X1)
    s3.shapes.add_picture(str(ROOT / "tmp.png"), Inches(1), Inches(1), Inches(2), Inches(2))
    ROOT.joinpath("tmp.png").unlink()
    p.save(ROOT / "deck.pptx")


if __name__ == "__main__":
    ROOT.mkdir(parents=True, exist_ok=True)
    build_docx()
    build_pptx()
    sys.stdout.write(f"fixtures written to {ROOT}\n")
